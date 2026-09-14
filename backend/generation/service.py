"""File generation service - creates real DOCX, PPTX, XLSX, PDF files."""
import os
import uuid
from datetime import datetime
from config import get_settings

settings = get_settings()


class FileGenerator:
    """Generates documents in various formats."""

    def __init__(self):
        os.makedirs(settings.GENERATED_DIR, exist_ok=True)

    async def generate(self, title: str, content: str, format: str, classification: str = "INTERNAL") -> str:
        file_id = str(uuid.uuid4())[:8]
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

        if format == "docx":
            return await self._generate_docx(title, content, file_id, timestamp, classification)
        elif format == "pptx":
            return await self._generate_pptx(title, content, file_id, timestamp, classification)
        elif format == "xlsx":
            return await self._generate_xlsx(title, content, file_id, timestamp, classification)
        elif format == "pdf":
            return await self._generate_pdf(title, content, file_id, timestamp, classification)
        return ""

    async def _generate_docx(self, title: str, content: str, file_id: str, timestamp: str, classification: str) -> str:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()

        # Header with classification
        header = doc.sections[0].header
        hp = header.paragraphs[0]
        hp.text = f"CLASSIFICATION: {classification}"
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Title
        title_para = doc.add_heading(title, level=0)

        # Metadata
        doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        doc.add_paragraph(f"Classification: {classification}")
        doc.add_paragraph("")

        # Content
        for paragraph in content.split("\n"):
            if paragraph.strip():
                doc.add_paragraph(paragraph.strip())

        # Footer
        footer = doc.sections[0].footer
        fp = footer.paragraphs[0]
        fp.text = f"Sovereign AI Workbench | {classification} | {file_id}"
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

        filename = f"{title.replace(' ', '_')}_{timestamp}.docx"
        filepath = os.path.join(settings.GENERATED_DIR, filename)
        doc.save(filepath)
        return filepath

    async def _generate_pptx(self, title: str, content: str, file_id: str, timestamp: str, classification: str) -> str:
        from pptx import Presentation
        from pptx.util import Inches, Pt

        prs = Presentation()

        # Title slide
        slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(slide_layout)
        slide.shapes.title.text = title
        slide.placeholders[1].text = f"Classification: {classification}\nGenerated: {datetime.now().strftime('%Y-%m-%d')}"

        # Content slides
        paragraphs = [p.strip() for p in content.split("\n") if p.strip()]
        for i in range(0, len(paragraphs), 5):
            chunk = paragraphs[i:i+5]
            slide_layout = prs.slide_layouts[1]
            slide = prs.slides.add_slide(slide_layout)
            slide.shapes.title.text = f"{title} ({i//5 + 1})"
            body = slide.placeholders[1]
            tf = body.text_frame
            tf.text = chunk[0]
            for line in chunk[1:]:
                p = tf.add_paragraph()
                p.text = line

        filename = f"{title.replace(' ', '_')}_{timestamp}.pptx"
        filepath = os.path.join(settings.GENERATED_DIR, filename)
        prs.save(filepath)
        return filepath

    async def _generate_xlsx(self, title: str, content: str, file_id: str, timestamp: str, classification: str) -> str:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = title[:31]

        # Header
        ws["A1"] = title
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = f"Classification: {classification}"
        ws["A3"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"

        # Content as rows
        row = 5
        for line in content.split("\n"):
            if line.strip():
                parts = line.split(",") if "," in line else [line]
                for col, val in enumerate(parts, 1):
                    ws.cell(row=row, column=col, value=val.strip())
                row += 1

        filename = f"{title.replace(' ', '_')}_{timestamp}.xlsx"
        filepath = os.path.join(settings.GENERATED_DIR, filename)
        wb.save(filepath)
        return filepath

    async def _generate_pdf(self, title: str, content: str, file_id: str, timestamp: str, classification: str) -> str:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.units import inch

        filename = f"{title.replace(' ', '_')}_{timestamp}.pdf"
        filepath = os.path.join(settings.GENERATED_DIR, filename)

        doc = SimpleDocTemplate(filepath, pagesize=A4)
        styles = getSampleStyleSheet()
        story = []

        # Classification header
        class_style = ParagraphStyle("Classification", parent=styles["Normal"], fontSize=8, alignment=1)
        story.append(Paragraph(f"CLASSIFICATION: {classification}", class_style))
        story.append(Spacer(1, 0.3 * inch))

        # Title
        story.append(Paragraph(title, styles["Title"]))
        story.append(Spacer(1, 0.2 * inch))

        # Metadata
        story.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", styles["Normal"]))
        story.append(Spacer(1, 0.3 * inch))

        # Content
        for para in content.split("\n"):
            if para.strip():
                story.append(Paragraph(para.strip(), styles["Normal"]))
                story.append(Spacer(1, 0.1 * inch))

        doc.build(story)
        return filepath
